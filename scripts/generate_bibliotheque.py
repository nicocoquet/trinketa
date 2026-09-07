#!/usr/bin/env python3
"""Construit la partie « Bibliothèque » du site MkDocs à partir d'Excel.

Entrées : ``inventaire_bibliotheque.xlsx`` (feuille ``Catalogue``) et les
originaux placés dans ``photos/bibliotheque/``.

Sorties : notices et index sous ``docs/biblioteca/``, copies d'images sous
``docs/assets/images/bibliotheque/`` et jeu de données
``docs/assets/data/bibliotheque-statistiques.json``.

Le JSON est l'interface entre Python et le navigateur : le générateur y place
les données nettoyées, puis ``bibliotheque-statistiques.js`` les charge pour
construire filtres, indicateurs, graphiques, classements et carte. Il ne doit
donc jamais être modifié à la main.

Ce script s'exécute après ``generate_mobilier.py`` dans le workflow.
"""

from __future__ import annotations

import html
import json
import re
import shutil
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "inventaire_bibliotheque.xlsx"
SOURCE_PHOTOS = ROOT / "photos" / "bibliotheque"
DOCS = ROOT / "docs"
LIBRARY = DOCS / "biblioteca"
LEGACY_LIBRARY = DOCS / "bibliotheque"
SITE_IMAGES = DOCS / "assets" / "images" / "bibliotheque"
SITE_IMAGES_URL = "/trinketa/assets/images/bibliotheque"
SITE_DATA = DOCS / "assets" / "data"

# Facettes retenues pour la bibliothèque Emeleta : pas de localisation,
# car l'immense majorité des ouvrages est conservée au Bureau.
FACETS = [
    ("author", "Auteur / éditeur scientifique"),
    ("publication_date", "Date de publication"),
    ("document_type", "Type de document"),
    ("subject", "Thème"),
    ("language", "Langue"),
]

STAT_FACETS = [
    ("publication_date", "Période"),
    ("language", "Langue"),
    ("subject", "Thème"),
    ("document_type", "Type de document"),
    ("heritage", "Intérêt patrimonial"),
]

# Repli prudent pour les villes les plus courantes. Les coordonnées du tableur
# restent prioritaires et permettent d'étendre la carte sans toucher au code.
KNOWN_CITIES = {
    "paris": (48.8566, 2.3522, "Paris", "France"),
    "bruyeres-le-chatel": (48.5935, 2.1925, "Bruyères-le-Châtel", "France"),
    "lyon": (45.7640, 4.8357, "Lyon", "France"),
    "strasbourg": (48.5734, 7.7521, "Strasbourg", "France"),
    "londres": (51.5074, -0.1278, "Londres", "Royaume-Uni"),
    "london": (51.5074, -0.1278, "Londres", "Royaume-Uni"),
    "rome": (41.9028, 12.4964, "Rome", "Italie"),
    "roma": (41.9028, 12.4964, "Rome", "Italie"),
    "milan": (45.4642, 9.1900, "Milan", "Italie"),
    "milano": (45.4642, 9.1900, "Milan", "Italie"),
    "florence": (43.7696, 11.2558, "Florence", "Italie"),
    "firenze": (43.7696, 11.2558, "Florence", "Italie"),
    "venise": (45.4408, 12.3155, "Venise", "Italie"),
    "venezia": (45.4408, 12.3155, "Venise", "Italie"),
    "turin": (45.0703, 7.6869, "Turin", "Italie"),
    "torino": (45.0703, 7.6869, "Turin", "Italie"),
    "zurich": (47.3769, 8.5417, "Zurich", "Suisse"),
    "geneve": (46.2044, 6.1432, "Genève", "Suisse"),
    "genève": (46.2044, 6.1432, "Genève", "Suisse"),
    "bruxelles": (50.8503, 4.3517, "Bruxelles", "Belgique"),
    "amsterdam": (52.3676, 4.9041, "Amsterdam", "Pays-Bas"),
    "berlin": (52.5200, 13.4050, "Berlin", "Allemagne"),
    "leipzig": (51.3397, 12.3731, "Leipzig", "Allemagne"),
    "vienne": (48.2082, 16.3738, "Vienne", "Autriche"),
    "wien": (48.2082, 16.3738, "Vienne", "Autriche"),
    "madrid": (40.4168, -3.7038, "Madrid", "Espagne"),
    "barcelone": (41.3874, 2.1686, "Barcelone", "Espagne"),
    "barcelona": (41.3874, 2.1686, "Barcelone", "Espagne"),
}

# Colonnes nécessaires au fonctionnement du site. Les colonnes d'enrichissement
# ISBN / notices sont optionnelles afin de ne pas bloquer la publication si elles
# sont progressivement ajoutées au tableur.
REQUIRED_COLUMNS = {
    "ID", "Auteur_editeur_scientifique", "Titre", "Date_publication", "Langue",
    "Type_document", "Sujet", "Mots_cles", "Serie_liee",
    "Photo_couverture", "Photo_bibliographique",
}


def text(value) -> str:
    """Convertit une cellule Excel en texte nettoyé, ou en chaîne vide."""
    return "" if value is None else str(value).strip()


def first_value(book: dict, *names: str) -> str:
    """Retourne la première valeur renseignée parmi plusieurs colonnes possibles."""
    for name in names:
        value = text(book.get(name))
        if value:
            return value
    return ""


def split_values(value) -> list[str]:
    """Découpe un champ multivalué séparé par des points-virgules."""
    return [item.strip() for item in text(value).split(";") if item.strip()]


def normalize_author(value: str) -> str:
    """Retire le rôle entre crochets pour stabiliser la facette d'auteur."""
    return re.sub(r"\s*\[[^\]]+\]\s*$", "", value).strip()


def display_author(value: str) -> str:
    """Met en forme les auteurs et leurs rôles pour l'affichage public."""
    authors = []
    for item in split_values(value):
        match = re.match(r"^(.*?)\s*\[([^\]]+)\]\s*$", item)
        if match:
            authors.append(f"{match.group(1).strip()} ({match.group(2).strip()})")
        else:
            authors.append(item)
    return " · ".join(authors)


def display_responsibilities(value: str) -> str:
    """Met en forme traducteurs, préfaciers, illustrateurs et autres rôles."""
    items = []
    for item in split_values(value):
        match = re.match(r"^(.*?)\s*\[([^\]]+)\]\s*$", item)
        if match:
            items.append(f"{match.group(1).strip()} ({match.group(2).strip()})")
        else:
            items.append(item)
    return " · ".join(items)


def author_facets(value: str) -> list[str]:
    """Retourne les noms normalisés utilisés par le filtre des auteurs."""
    return [normalize_author(item) for item in split_values(value)]


def publication_bucket(value) -> str:
    """Classe une date libre dans une période utilisée par les statistiques."""
    match = re.search(r"\b(\d{4})\b", text(value))
    if not match:
        return "Date à préciser"
    year = int(match.group(1))
    if year < 1900:
        return "Avant 1900"
    if year <= 1945:
        return "1900–1945"
    if year <= 1980:
        return "1946–1980"
    if year <= 2000:
        return "1981–2000"
    return "Après 2000"


def publication_year(value):
    """Extrait une année exploitable d'une date libre, ou renvoie ``None``."""
    match = re.search(r"\b(1[0-9]{3}|20[0-9]{2})\b", text(value))
    return int(match.group(1)) if match else None


def number(value):
    """Convertit une cellule numérique, y compris avec une virgule décimale."""
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def normalized_key(value: str) -> str:
    """Normalise casse et accents pour comparer des libellés de façon robuste."""
    value = unicodedata.normalize("NFD", text(value).casefold())
    return "".join(ch for ch in value if unicodedata.category(ch) != "Mn").strip()


def location_data(book: dict):
    """Résout ville, pays et coordonnées, avec repli sur ``KNOWN_CITIES``."""
    city = first_value(book, "Ville_normalisee", "Lieu_publication")
    country = text(book.get("Pays_normalise"))
    latitude = number(book.get("Latitude"))
    longitude = number(book.get("Longitude"))
    fallback = KNOWN_CITIES.get(normalized_key(city))
    if fallback:
        latitude = latitude if latitude is not None else fallback[0]
        longitude = longitude if longitude is not None else fallback[1]
        city = first_value(book, "Ville_normalisee") or fallback[2]
        country = country or fallback[3]
    return city, country, latitude, longitude


def truthy(value) -> bool:
    """Reconnaît les principales écritures de « oui » dans les trois langues."""
    return normalized_key(text(value)) in {"oui", "yes", "si", "true", "1"}


def slug_sort(value: str) -> str:
    """Produit une clé de tri insensible aux accents et à la casse."""
    normalized = unicodedata.normalize("NFD", value)
    return "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn").casefold()


def rows_as_dicts(sheet):
    """Transforme chaque ligne Excel non vide en dictionnaire nommé par colonne."""
    headers = [text(cell.value) for cell in sheet[1]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue
        yield {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}


def require_columns(sheet):
    """Interrompt la génération si une colonne indispensable a été renommée."""
    actual = {text(cell.value) for cell in sheet[1]}
    missing = REQUIRED_COLUMNS - actual
    if missing:
        raise ValueError("Colonnes manquantes dans Catalogue : " + ", ".join(sorted(missing)))


def data_attr(values) -> str:
    """Encode plusieurs facettes dans un attribut HTML séparé par ``||``."""
    return "||".join(values)


def copy_photo(filename: str, warnings: list[str]) -> str:
    """Copie une photo vers MkDocs et signale une absence sans bloquer le site."""
    filename = Path(filename).name
    if not filename:
        return ""
    source = SOURCE_PHOTOS / filename
    if not source.is_file():
        warnings.append(f"Photographie introuvable : {filename}")
        return ""
    SITE_IMAGES.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, SITE_IMAGES / filename)
    return filename


def field(label: str, value: str) -> str:
    """Construit une paire libellé/valeur HTML si elle est renseignée."""
    if not value:
        return ""
    return f'<div class="record-field"><dt>{html.escape(label)}</dt><dd>{html.escape(value)}</dd></div>'


def metadata_html(book: dict) -> str:
    """Construit le bloc de description bibliographique d'une notice."""
    publication = ", ".join(
        part for part in [
            text(book.get("Lieu_publication")),
            text(book.get("Editeur")),
            text(book.get("Date_publication")),
        ] if part
    )
    fields = [
        ("Auteur / éditeur scientifique", display_author(text(book.get("Auteur_editeur_scientifique")))),
        ("Autre responsabilité", display_responsibilities(first_value(book, "Autre_responsabilite", "Autre_responsabilité"))),
        ("Sous-titre", text(book.get("Sous_titre"))),
        ("Volume", text(book.get("Volume"))),
        ("Publication", publication),
        ("Édition", text(book.get("Edition"))),
        ("Collection", text(book.get("Collection"))),
        ("Numéro dans la collection", first_value(book, "Num_collection", "Numero_collection")),
        ("ISBN", first_value(book, "ISBN", "ISBN_13")),
        ("ISBN-13", first_value(book, "ISBN_13", "EAN")),
        ("ISBN-10", text(book.get("ISBN_10"))),
        ("Langue", text(book.get("Langue"))),
        ("Langue originale", text(book.get("Langue_originale"))),
        ("Titre original", text(book.get("Titre_original"))),
        ("Type de document", text(book.get("Type_document"))),
        ("Genre", text(book.get("Genre"))),
        ("Nombre de pages", text(book.get("Nbr_pages"))),
        ("Dimensions", text(book.get("Dimensions"))),
        ("Localisation", text(book.get("Localisation"))),
    ]
    return "\n".join(field(label, value) for label, value in fields if value)


def indexing_html(book: dict) -> str:
    """Construit le bloc d'indexation thématique et chronologique."""
    fields = [
        ("Thème", text(book.get("Sujet"))),
        ("Lieu", text(book.get("Lieu_sujet"))),
        ("Période", text(book.get("Periode_sujet"))),
        ("Mots-clés", text(book.get("Mots_cles"))),
    ]
    return "\n".join(field(label, value) for label, value in fields if value)


def notice_html(book: dict) -> str:
    """Construit les informations de provenance et de contrôle de la notice."""
    fields = [
        ("Source de la notice", text(book.get("Source_notice"))),
        ("Identifiant de la notice", text(book.get("Identifiant_notice"))),
        ("Statut de la notice", text(book.get("Statut_notice"))),
        ("Confiance", text(book.get("Confiance_notice"))),
        ("Notes de contrôle", text(book.get("Notes_controle"))),
    ]
    return "\n".join(field(label, value) for label, value in fields if value)


def exemplar_html(book: dict) -> str:
    """Construit les particularités propres à l'exemplaire inventorié."""
    fields = [
        ("Ex-libris", text(book.get("Ex_libris"))),
        ("Dédicace / annotations", text(book.get("Dedicace_annotations"))),
        ("Notes", text(book.get("Notes"))),
    ]
    return "\n".join(field(label, value) for label, value in fields if value)


def render_select(name: str, label: str, values: list[str]) -> str:
    """Construit une liste déroulante de facette triée et échappée."""
    options = ['<option value="">Tous</option>']
    options.extend(
        f'<option value="{html.escape(value, quote=True)}">{html.escape(value)}</option>'
        for value in sorted(values, key=slug_sort)
    )
    return (
        f'<div class="catalog-filter"><label for="filter-{name}">{html.escape(label)}</label>'
        f'<select id="filter-{name}" name="{name}" data-filter="{name}">' + "".join(options) + "</select></div>"
    )



def main() -> int:
    """Orchestre notices, index, images et jeu de données statistiques JSON."""
    if not WORKBOOK.exists():
        print(f"Classeur introuvable : {WORKBOOK}", file=sys.stderr)
        return 1
    wb = load_workbook(WORKBOOK, data_only=True)
    if "Catalogue" not in wb.sheetnames:
        print("Feuille Catalogue introuvable.", file=sys.stderr)
        return 1
    sheet = wb["Catalogue"]
    require_columns(sheet)
    books = [row for row in rows_as_dicts(sheet) if text(row.get("ID")) and text(row.get("Titre"))]
    known_ids = {text(book["ID"]).upper() for book in books}
    LIBRARY.mkdir(parents=True, exist_ok=True)
    SITE_IMAGES.mkdir(parents=True, exist_ok=True)
    if LEGACY_LIBRARY.exists():
        for legacy in [*LEGACY_LIBRARY.glob("BIB-*.md"), *LEGACY_LIBRARY.glob("index*.md"), *LEGACY_LIBRARY.glob("statistiques.*.md")]:
            legacy.unlink()
        try:
            LEGACY_LIBRARY.rmdir()
        except OSError:
            # Ne pas supprimer un dossier contenant un fichier manuel inconnu.
            pass
    for legacy in LIBRARY.glob("BIB-*.fr.md"):
        legacy.unlink()
    facet_values = {name: set() for name, _ in FACETS}
    cards = []
    statistics = []
    warnings: list[str] = []
    for book in books:
        book_id = text(book["ID"]).upper()
        title = text(book["Titre"])
        authors = author_facets(text(book.get("Auteur_editeur_scientifique")))
        subjects = split_values(book.get("Sujet"))
        publication_group = publication_bucket(book.get("Date_publication"))
        languages = split_values(book.get("Langue"))
        document_types = split_values(book.get("Type_document"))
        per_book_facets = {
            "author": authors,
            "publication_date": [publication_group],
            "document_type": document_types,
            "subject": subjects,
            "language": languages,
        }
        city, country, latitude, longitude = location_data(book)
        low = number(book.get("Estimation_basse_EUR"))
        high = number(book.get("Estimation_haute_EUR"))
        volumes = number(book.get("Nombre_volumes")) or 1
        statistics.append({
            "id": book_id,
            "title": title,
            "year": publication_year(book.get("Date_publication")),
            "period": publication_group,
            "languages": languages,
            "subjects": subjects,
            "publishers": split_values(book.get("Editeur")),
            "documentTypes": document_types,
            "heritage": text(book.get("Interet_patrimonial")) or "À préciser",
            "completeness": text(book.get("Completude")) or "À préciser",
            "illustrated": text(book.get("Illustre")) or "À préciser",
            "marks": split_values(book.get("Marques_exemplaire")),
            "volumes": volumes,
            "estimateLow": low,
            "estimateHigh": high,
            "cityHistorical": text(book.get("Lieu_publication")),
            "city": city,
            "country": country,
            "latitude": latitude,
            "longitude": longitude,
            "locationUncertain": truthy(book.get("Lieu_incertitude")),
        })
        for name, values in per_book_facets.items():
            facet_values[name].update(values)
        cover = copy_photo(text(book.get("Photo_couverture")), warnings)
        biblio_photo = copy_photo(text(book.get("Photo_bibliographique")), warnings)
        extra_photos = [
            copied for item in split_values(book.get("Photos_supplementaires"))
            if (copied := copy_photo(item, warnings))
        ]
        cover_html = (
            f'<a href="{SITE_IMAGES_URL}/{html.escape(cover)}" target="_blank">'
            f'<img src="{SITE_IMAGES_URL}/{html.escape(cover)}" alt="Couverture — {html.escape(title)}"></a>'
            if cover else '<div class="record-placeholder book-placeholder">Sans photographie</div>'
        )
        gallery_items = []
        if biblio_photo:
            gallery_items.append(
                f'<figure><a href="{SITE_IMAGES_URL}/{html.escape(biblio_photo)}" target="_blank">'
                f'<img src="{SITE_IMAGES_URL}/{html.escape(biblio_photo)}" alt="Informations bibliographiques — {html.escape(title)}" loading="lazy"></a>'
                f'<figcaption>Page de titre / informations bibliographiques</figcaption></figure>'
            )
        for filename in extra_photos:
            gallery_items.append(
                f'<figure><a href="{SITE_IMAGES_URL}/{html.escape(filename)}" target="_blank">'
                f'<img src="{SITE_IMAGES_URL}/{html.escape(filename)}" alt="{html.escape(title)}" loading="lazy"></a>'
                f'<figcaption>Photographie complémentaire</figcaption></figure>'
            )
        gallery = "\n".join(gallery_items) or '<p class="empty-state">Aucune photographie complémentaire.</p>'
        links = []
        for associated in re.findall(r"BIB-\d{3,}", text(book.get("Serie_liee")), flags=re.IGNORECASE):
            associated = associated.upper()
            if associated == book_id:
                continue
            if associated in known_ids:
                links.append(f'<a href="../{associated}/">{html.escape(associated)}</a>')
            else:
                warnings.append(f"{book_id} : volume lié inconnu : {associated}")
        related = " · ".join(links) or "_Aucun volume associé._"
        external = first_value(book, "URL_notice", "URL_source")
        external_block = f'[{html.escape(external)}]({external})' if external else "_Aucune notice externe renseignée._"
        notice = notice_html(book)
        page = f'''<div class="lot-hero book-hero">\n<div class="lot-visual book-cover">{cover_html}</div>\n<div class="lot-summary">\n<p class="record-kicker">Bibliothèque · {html.escape(book_id)}</p>\n<h1>{html.escape(title)}</h1>\n<p class="lot-dating">{html.escape(display_author(text(book.get("Auteur_editeur_scientifique"))))}</p>\n<p class="lot-estimate">{html.escape(", ".join(part for part in [text(book.get("Lieu_publication")), text(book.get("Editeur")), text(book.get("Date_publication"))] if part))}</p>\n</div>\n</div>\n\n## Notice bibliographique\n\n<dl class="record-metadata">\n{metadata_html(book)}\n</dl>\n\n## Indexation\n\n<dl class="record-metadata">\n{indexing_html(book)}\n</dl>\n\n## Données de notice\n\n<dl class="record-metadata">\n{notice or '<div class="record-field"><dd>Aucune donnée de notice renseignée.</dd></div>'}\n</dl>\n\n## Exemplaire d'Emeleta\n\n<dl class="record-metadata">\n{exemplar_html(book) or '<div class="record-field"><dd>Aucune particularité renseignée.</dd></div>'}\n</dl>\n\n## Volumes associés\n\n{related}\n\n## Notice externe\n\n{external_block}\n\n## Photographies\n\n<div class="photo-grid book-photo-grid">\n{gallery}\n</div>\n\n[← Retour à la bibliothèque](../)\n'''
        (LIBRARY / f"{book_id}.fr.md").write_text(page, encoding="utf-8")
        attrs = " ".join(
            f'data-{name.replace("_", "-" )}="{html.escape(data_attr(values), quote=True)}"'
            for name, values in per_book_facets.items()
        )
        visual = (
            f'<img src="{SITE_IMAGES_URL}/{html.escape(cover)}" alt="{html.escape(title)}" loading="lazy">'
            if cover else '<div class="card-placeholder book-placeholder">Sans photographie</div>'
        )
        author_label = " · ".join(authors) or "Auteur non renseigné"
        pub_label = ", ".join(part for part in [text(book.get("Editeur")), text(book.get("Date_publication"))] if part)
        card_text = " | ".join(
            part for part in [title, author_label, pub_label, text(book.get("Mots_cles")), text(book.get("ISBN")), text(book.get("ISBN_13"))]
            if part
        )
        cards.append(
            f'<article class="catalog-card book-card" {attrs} data-search="{html.escape(card_text, quote=True)}">'
            f'<a href="{html.escape(book_id)}/">{visual}</a>'
            f'<div class="catalog-card-body"><div class="card-lot"><span>{html.escape(book_id)}</span><span>{html.escape(text(book.get("Type_document")))}</span></div>'
            f'<h2><a href="{html.escape(book_id)}/">{html.escape(title)}</a></h2>'
            f'<p class="book-author">{html.escape(author_label)}</p><p class="card-date">{html.escape(pub_label)}</p></div></article>'
        )
    selects = "\n".join(render_select(name, label, list(facet_values[name])) for name, label in FACETS)
    index = f'''<div class="catalogue-heading">\n<h1>Bibliothèque</h1>\n<a class="library-isbn-entry" href="/trinketa/depot-isbn/" target="_blank" rel="noopener">Déposer des photographies ISBN <span>→</span></a>\n</div>\n\n<form class="catalog-filters library-filters" data-catalog-filters>\n<div class="catalog-filter catalog-filter-search"><label for="catalog-search">Recherche</label><input id="catalog-search" type="search" name="q" placeholder="Titre, auteur, ISBN, mot-clé…" autocomplete="off"></div>\n{selects}\n<button class="catalog-reset" type="reset">Réinitialiser</button>\n<p class="catalog-result"><strong data-result-count>{len(cards)}</strong> ouvrage(s)</p>\n</form>\n\n<div class="catalog-grid book-grid" data-catalog-grid>\n{chr(10).join(cards)}\n</div>\n\n<p class="catalog-empty" data-catalog-empty hidden>Aucun ouvrage ne correspond à ces critères.</p>\n'''
    (LIBRARY / "index.fr.md").write_text(index, encoding="utf-8")
    SITE_DATA.mkdir(parents=True, exist_ok=True)
    data = {
        "generatedFrom": WORKBOOK.name,
        "records": statistics,
        "facets": [name for name, _ in STAT_FACETS],
    }
    (SITE_DATA / "bibliotheque-statistiques.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    # La page Statistiques est désormais unifiée à la racine du site et
    # générée par generate_mobilier.py. Supprimer les anciennes pages séparées
    # évite qu'elles restent accessibles par la recherche ou des liens directs.
    for language in ("fr", "en", "it"):
        legacy_statistics = LIBRARY / f"statistiques.{language}.md"
        if legacy_statistics.exists():
            legacy_statistics.unlink()
    print(f"{len(cards)} ouvrage(s) généré(s).")
    if warnings:
        print("Avertissements :", file=sys.stderr)
        for warning in sorted(set(warnings)):
            print(f"- {warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
